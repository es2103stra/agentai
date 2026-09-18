"""
Document Agent Examples
========================
Примеры использования модуля для работы с документами.
"""

from document_agent import DocumentAnalyzer, DocumentGenerator, AutoReportGenerator
import json

def example_1_analyze_document():
    """Пример 1: Анализ существующего документа."""
    print("\n" + "="*70)
    print("📄 Пример 1: Анализ документа")
    print("="*70)
    
    # Создаём тестовый документ
    from docx import Document
    doc = Document()
    doc.add_heading('Квартальный отчёт Q4 2024', 0)
    doc.add_heading('Введение', 1)
    doc.add_paragraph('Этот отчёт содержит данные за {{quarter}} квартал {{year}} года.')
    doc.add_heading('Основные показатели', 1)
    doc.add_paragraph('Выручка: {{revenue}} рублей')
    doc.add_paragraph('Прибыль: {{profit}} рублей')
    doc.add_heading('Заключение', 1)
    doc.add_paragraph('{{conclusion}}')
    doc.save('test_report.docx')
    
    # Анализируем
    analyzer = DocumentAnalyzer()
    structure = analyzer.analyze('test_report.docx')
    
    print(f"\n📊 Структура:")
    print(f"   Заголовок: {structure.title}")
    print(f"   Секций: {len(structure.sections)}")
    print(f"   Плейсхолдеров: {analyzer.extract_placeholders()}")
    
    # Экспортируем в JSON
    analyzer.export_structure('test_report_structure.json')
    print(f"\n✅ Структура сохранена в test_report_structure.json")


def example_2_generate_from_data():
    """Пример 2: Генерация документа из данных."""
    print("\n" + "="*70)
    print("📝 Пример 2: Генерация из данных")
    print("="*70)
    
    # Создаём шаблон
    from docx import Document
    doc = Document()
    doc.add_heading('Отчёт по проекту {{project_name}}', 0)
    doc.add_heading('Описание', 1)
    doc.add_paragraph('{{project_description}}')
    doc.add_heading('Результаты', 1)
    doc.add_paragraph('Срок выполнения: {{deadline}}')
    doc.add_paragraph('Бюджет: {{budget}} рублей')
    doc.add_heading('Выводы', 1)
    doc.add_paragraph('{{conclusions}}')
    doc.save('project_template.docx')
    
    # Данные для заполнения
    data = {
        "project_name": "Автоматизация отчётности",
        "project_description": "Разработка системы автоматической генерации отчётов на основе шаблонов DOCX.",
        "deadline": "31 декабря 2024",
        "budget": "500,000",
        "conclusions": "Проект успешно завершён в срок. Все цели достигнуты."
    }
    
    # Генерируем документ
    generator = DocumentGenerator()
    generator.load_template('project_template.docx')
    result = generator.generate_from_data(data, 'project_report_generated.docx')
    
    print(f"\n✅ Документ создан: {result.path}")
    print(f"   Слов: {result.word_count}")


def example_3_auto_report_with_ai():
    """Пример 3: Автоматическая генерация отчёта с AI."""
    print("\n" + "="*70)
    print("🤖 Пример 3: Автогенерация с AI")
    print("="*70)
    
    # Создаём шаблон с плейсхолдерами
    from docx import Document
    doc = Document()
    doc.add_heading('Маркетинговый анализ', 0)
    doc.add_heading('Обзор рынка', 1)
    doc.add_paragraph('{{market_overview}}')
    doc.add_heading('Конкуренты', 1)
    doc.add_paragraph('{{competitors_analysis}}')
    doc.add_heading('Рекомендации', 1)
    doc.add_paragraph('{{recommendations}}')
    doc.save('marketing_template.docx')
    
    # Генерируем с AI
    auto_gen = AutoReportGenerator()
    result = auto_gen.generate_report(
        template_path='marketing_template.docx',
        task='Создай маркетинговый анализ для запуска нового мобильного приложения в России',
        output_path='marketing_report_ai.docx'
    )
    
    print(f"\n✅ Отчёт создан: {result.path}")
    print(f"   Плейсхолдеров заполнено: {result.placeholders_filled}")


def example_4_batch_reports():
    """Пример 4: Пакетная генерация отчётов."""
    print("\n" + "="*70)
    print("📦 Пример 4: Пакетная генерация")
    print("="*70)
    
    # Создаём шаблон
    from docx import Document
    doc = Document()
    doc.add_heading('Еженедельный отчёт', 0)
    doc.add_heading('Достижения', 1)
    doc.add_paragraph('{{achievements}}')
    doc.add_heading('Проблемы', 1)
    doc.add_paragraph('{{issues}}')
    doc.add_heading('Планы', 1)
    doc.add_paragraph('{{plans}}')
    doc.save('weekly_template.docx')
    
    # Задачи для генерации
    tasks = [
        "Отчёт за неделю 1: Запуск нового функционала",
        "Отчёт за неделю 2: Исправление багов",
        "Отчёт за неделю 3: Оптимизация производительности"
    ]
    
    # Генерируем пакетно
    auto_gen = AutoReportGenerator()
    results = auto_gen.batch_generate('weekly_template.docx', tasks)
    
    print(f"\n✅ Создано {len(results)} отчётов:")
    for i, result in enumerate(results, 1):
        print(f"   {i}. {result.path}")


def example_5_complex_template():
    """Пример 5: Сложный шаблон с таблицами."""
    print("\n" + "="*70)
    print("📊 Пример 5: Сложный шаблон")
    print("="*70)
    
    # Создаём сложный шаблон
    from docx import Document
    doc = Document()
    doc.add_heading('Финансовый отчёт {{company_name}}', 0)
    
    doc.add_heading('Резюме', 1)
    doc.add_paragraph('{{executive_summary}}')
    
    doc.add_heading('Финансовые показатели', 1)
    table = doc.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    
    # Заголовки таблицы
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Показатель'
    hdr_cells[1].text = 'Q3 2024'
    hdr_cells[2].text = 'Q4 2024'
    
    # Данные (с плейсхолдерами)
    row1 = table.rows[1].cells
    row1[0].text = 'Выручка'
    row1[1].text = '{{q3_revenue}}'
    row1[2].text = '{{q4_revenue}}'
    
    row2 = table.rows[2].cells
    row2[0].text = 'Расходы'
    row2[1].text = '{{q3_expenses}}'
    row2[2].text = '{{q4_expenses}}'
    
    row3 = table.rows[3].cells
    row3[0].text = 'Прибыль'
    row3[1].text = '{{q3_profit}}'
    row3[2].text = '{{q4_profit}}'
    
    doc.add_heading('Анализ', 1)
    doc.add_paragraph('{{financial_analysis}}')
    
    doc.save('financial_template.docx')
    
    # Анализируем
    analyzer = DocumentAnalyzer()
    structure = analyzer.analyze('financial_template.docx')
    
    print(f"\n📊 Структура:")
    print(f"   Таблиц: {structure.tables_count}")
    print(f"   Плейсхолдеров: {len(analyzer.extract_placeholders())}")
    
    # Извлекаем таблицы
    tables = analyzer.extract_tables()
    print(f"\n📋 Таблица 1:")
    for row in tables[0]:
        print(f"   {row}")


def example_6_custom_styling():
    """Пример 6: Документ с кастомными стилями."""
    print("\n" + "="*70)
    print("🎨 Пример 6: Кастомные стили")
    print("="*70)
    
    from docx import Document
    from docx.shared import Pt, RGBColor
    
    doc = Document()
    
    # Заголовок с кастомным стилем
    title = doc.add_heading('{{report_title}}', 0)
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.color.rgb = RGBColor(0, 51, 102)
    
    # Подзаголовок
    subtitle = doc.add_heading('{{subtitle}}', 1)
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(102, 102, 102)
    
    # Основной текст
    doc.add_paragraph('{{main_content}}')
    
    # Выделенный блок
    highlight = doc.add_paragraph()
    run = highlight.add_run('{{highlight_text}}')
    run.bold = True
    run.font.color.rgb = RGBColor(204, 0, 0)
    
    doc.save('styled_template.docx')
    
    # Анализируем стили
    analyzer = DocumentAnalyzer()
    structure = analyzer.analyze('styled_template.docx')
    
    print(f"\n🎨 Стили:")
    for style in structure.styles_used:
        print(f"   - {style}")
    
    # Генерируем документ
    data = {
        "report_title": "Годовой отчёт 2024",
        "subtitle": "Компания XYZ",
        "main_content": "Это основной текст отчёта с подробным анализом...",
        "highlight_text": "ВАЖНО: Прибыль выросла на 25%!"
    }
    
    generator = DocumentGenerator()
    generator.load_template('styled_template.docx')
    result = generator.generate_from_data(data, 'styled_report.docx')
    
    print(f"\n✅ Документ создан: {result.path}")


if __name__ == "__main__":
    import sys
    
    examples = {
        "1": ("Анализ документа", example_1_analyze_document),
        "2": ("Генерация из данных", example_2_generate_from_data),
        "3": ("Автогенерация с AI", example_3_auto_report_with_ai),
        "4": ("Пакетная генерация", example_4_batch_reports),
        "5": ("Сложный шаблон", example_5_complex_template),
        "6": ("Кастомные стили", example_6_custom_styling),
    }
    
    print("\n" + "="*70)
    print("📄 Document Agent — Examples")
    print("="*70)
    print("\nВыберите пример:")
    
    for key, (name, _) in examples.items():
        print(f"  {key}. {name}")
    
    print("\n  0. Выход")
    
    choice = input("\nВаш выбор: ").strip()
    
    if choice == "0":
        print("👋 До свидания!")
        sys.exit(0)
    
    if choice in examples:
        name, func = examples[choice]
        print(f"\n🚀 Запускаю: {name}\n")
        func()
    else:
        print("❌ Неверный выбор")
